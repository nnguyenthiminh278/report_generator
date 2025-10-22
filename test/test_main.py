import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import sqlite3
import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm  # for image sizing
import os
import glob
from datetime import datetime
from docx import Document

class ReportGeneratorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Report Generator")

        # Search type dropdown
        tk.Label(root, text="Search by:").grid(row=0, column=0, padx=5, pady=5)
        self.search_type = ttk.Combobox(root, values=["Patient ID", "Sample ID", "Name"])
        self.search_type.current(0)  # default to Patient ID
        self.search_type.grid(row=0, column=1, padx=5, pady=5)

        # Search value entry
        tk.Label(root, text="Value:").grid(row=1, column=0, padx=5, pady=5)
        self.search_value_entry = tk.Entry(root)
        self.search_value_entry.grid(row=1, column=1, padx=5, pady=5)

        # Working directory selection
        tk.Button(root, text="Select Directory", command=self.select_directory).grid(row=2, column=0, padx=5, pady=5)
        self.dir_label = tk.Label(root, text="No directory selected")
        self.dir_label.grid(row=2, column=1, padx=5, pady=5)

        # Generate button
        tk.Button(root, text="Generate Report", command=self.generate_report).grid(row=3, column=0, columnspan=2, pady=10)

        # Button for annex report
        tk.Button(root, text="Generate Annex", command=self.generate_annex).grid(row=4, column=0, columnspan=2, pady=10)

    def select_directory(self):
        self.working_dir = filedialog.askdirectory()
        if self.working_dir:
            self.dir_label.config(text=self.working_dir)

    def generate_report(self):
        search_type = self.search_type.get()
        search_value = self.search_value_entry.get().strip()

        if not search_value:
            messagebox.showerror("Error", "Please enter a search value")
            return

        # --- DB ---
        conn = sqlite3.connect("patients.db")
        cur = conn.cursor()

        query = """
            SELECT first_name, last_name, dob, gender, patient_id, sample_id,
                   analysis_id, sample_date, address, diagnosis
            FROM patients
        """
        if search_type == "Patient ID":
            query += " WHERE patient_id=?"
            cur.execute(query, (search_value,))
        elif search_type == "Sample ID":
            query += " WHERE sample_id=?"
            cur.execute(query, (search_value,))
        elif search_type == "Name":
            query += " WHERE last_name=? OR first_name=?"
            cur.execute(query, (search_value, search_value))

        result = cur.fetchone()
        conn.close()

        if not result:
            messagebox.showerror("Error", f"No patient found with {search_type} '{search_value}'")
            return

        (first_name, last_name, dob, gender,
         patient_id_str, sample_id, analysis_id,
         sample_date, address, diagnosis) = result

        # Normalize gender and decide salutation
        if str(gender).lower().strip() in ["weiblich", "female", "f"]:
            anrede = "Frau"
        else:
            anrede = "Herr"

        # --- Excel ---
        excel_path = os.path.join(self.working_dir, "Mustertabelle_Klassifikation_alles.xlsx")
        df = pd.read_excel(excel_path, header=None)

        # Row 3 (index 2) has headers
        headers = df.iloc[2].tolist()

        # Find row with "final score" (usually in column 2 / index 1)
        final_row = df[df.iloc[:, 1].astype(str).str.lower().str.contains("final score")]

        if final_row.empty:
            raise ValueError("No 'final score' row found in Excel")

        scores_dict = dict(zip(headers, final_row.iloc[0].tolist()))

        # Mapping: Excel column → report placeholder
        mapping = {
            "CAD238ML1k.mdl": "CAD_score",
            "CKD273ML1hybrid": "CKD_score",
            "HF2_ML1new.mdl": "HF_score",
            "oncoRisk normo": "Onkorisk_score",
            "BioAge": "BioAge_value"
        }

        mapped_scores = {
            v: f"{scores_dict[k]:.3f}"  # 3 decimal rounding
            for k, v in mapping.items()
            if k in scores_dict
        }

        # --- Read thresholds and compare to scores ---
        doc = Document("./templates/template_MOS.docx")
        target_table = None
        for tbl in doc.tables:
            header = [cell.text.strip() for cell in tbl.rows[0].cells]
            if "Normalbereich" in header:
                target_table = tbl
                break

        if target_table is None:
            raise ValueError("No table with 'Normalbereich' column found!")
        # Mapping: table row index → score key
        row_mapping = {
            1: "CKD_score",       # 2nd row in table (KidneyRisk)
            2: "CAD_score",       # 3rd row (HeartRisk)
            3: "HF_score",        # 4th row (Herzinsuffizienz)
            4: "Onkorisk_score"   # 5th row (OncoRisk)
        }

        sentences = {}

        for row_idx, score_key in row_mapping.items():
            # Normalbereich column = 3 (4th column)
            normalbereich_text = target_table.cell(row_idx, 3).text.strip()
            
            try:
                threshold = float(normalbereich_text.replace("<", "").strip().replace(",", "."))
            except ValueError:
                threshold = None

            value_str = mapped_scores.get(score_key, "0").replace(",", ".")
            value = float(value_str)

            if threshold is not None:
                if value < threshold:
                    sentences[score_key] = "keine"
                else:
                    sentences[score_key] = "eine"
            else:
                sentences[score_key] = "keine"  # fallback

        # --- Word template with docxtpl ---
        tpl = DocxTemplate("./templates/template_MOS.docx")

        # Report generation date (German format: DD.MM.YYYY)
        report_date = datetime.now().strftime("%d.%m.%Y")

        context = {
            "anrede": anrede,
            "vorname": first_name,
            "name": last_name,
            "dob": dob,
            "geschlecht": gender,
            "patient_id": patient_id_str,
            "sample_id": sample_id,
            "analysis_id": analysis_id,
            "sample_date": sample_date,
            "address": address,
            "diagnosis": diagnosis,
            "report_date": report_date,
            **mapped_scores
        }

        # Add into context for docxtpl
        context.update({
            "ckd_sentence": sentences["CKD_score"],
            "cad_sentence": sentences["CAD_score"],
            "hf_sentence": sentences["HF_score"],
            "onco_sentence": sentences["Onkorisk_score"]
        })

        # Insert figures (based on postfix like *_1.png → {{ fig1 }})
        for file in glob.glob(os.path.join(self.working_dir, "*_[0-9].png")):
            filename = os.path.basename(file)
            number = filename.split("_")[-1].replace(".png", "")
            placeholder = f"fig{number}"
            context[placeholder] = InlineImage(tpl, file, width=Mm(80))

        tpl.render(context)

        output_path = os.path.join(self.working_dir, f"{last_name}_{first_name}_Report.docx")
        tpl.save(output_path)

        messagebox.showinfo("Success", f"Report saved as {output_path}")

    def generate_annex(self):
        search_type = self.search_type.get()
        search_value = self.search_value_entry.get().strip()

        if not search_value:
            messagebox.showerror("Error", "Please enter a search value")
            return

        # --- DB lookup (reuse from main report) ---
        conn = sqlite3.connect("patients.db")
        cur = conn.cursor()

        query = """
            SELECT first_name, last_name, dob, gender, patient_id, sample_id,
                analysis_id, sample_date, address, diagnosis
            FROM patients
        """
        if search_type == "Patient ID":
            query += " WHERE patient_id=?"
            cur.execute(query, (search_value,))
        elif search_type == "Sample ID":
            query += " WHERE sample_id=?"
            cur.execute(query, (search_value,))
        elif search_type == "Name":
            query += " WHERE last_name=? OR first_name=?"
            cur.execute(query, (search_value, search_value))

        result = cur.fetchone()
        conn.close()

        if not result:
            messagebox.showerror("Error", f"No patient found with {search_type} '{search_value}'")
            return

        (first_name, last_name, dob, gender,
        patient_id_str, sample_id, analysis_id,
        sample_date, address, diagnosis) = result

        # Salutation
        if str(gender).lower().strip() in ["weiblich", "female", "f"]:
            anrede = "Frau"
        else:
            anrede = "Herr"

        # Report date
        report_date = datetime.now().strftime("%d.%m.%Y")

        # Base context
        context = {
            "anrede": anrede,
            "vorname": first_name,
            "name": last_name,
            "dob": dob,
            "geschlecht": gender,
            "patient_id": patient_id_str,
            "sample_id": sample_id,
            "analysis_id": analysis_id,
            "sample_date": sample_date,
            "address": address,
            "diagnosis": diagnosis,
            "report_date": report_date,
            "patient_sample": f"{patient_id_str}-{sample_id}"   # special combined field
        }

        # --- Load annex template ---
        tpl = DocxTemplate("./templates/template_annex.docx")

        # Insert annex figures (e.g., *_1.png → {{ annex_fig1 }})
        for file in glob.glob(os.path.join(self.working_dir, "*_[0-9].png")):
            filename = os.path.basename(file)
            number = filename.split("_")[-1].replace(".png", "")
            placeholder = f"annex_fig{number}"   # e.g., annex_fig1
            context[placeholder] = InlineImage(tpl, file, width=Mm(80))

        # --- Render and save annex ---
        tpl.render(context)

        output_path = os.path.join(self.working_dir, f"{last_name}_{first_name}_Annex.docx")
        tpl.save(output_path)

        messagebox.showinfo("Success", f"Annex saved as {output_path}")



if __name__ == "__main__":
    root = tk.Tk()
    app = ReportGeneratorApp(root)
    root.mainloop()
